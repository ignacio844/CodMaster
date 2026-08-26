import json
from pathlib import Path

import openpyxl

SOURCE = Path(r"C:\Users\Auditoria\Downloads\2026.08.25 15.04 - Productos.xlsx")
OUTPUT = Path(r"C:\Users\Auditoria\Documents\Codex\2026-08-25\database-cod-sku\work\app\data\bam-master.json")

workbook = openpyxl.load_workbook(SOURCE, read_only=True, data_only=True)
sheet = workbook["Productos"]
rows = sheet.iter_rows(values_only=True)
headers = [str(value or "").strip() for value in next(rows)]
index = {header: position for position, header in enumerate(headers)}

products = {}
for row in rows:
    code = str(row[index["CODIGO"]] or "").strip()
    if not code:
        continue
    current = products.get(code)
    if current is None:
        current = {
            "code": code,
            "alternativeCode": str(row[index["CODIGO ALTERNATIVO"]] or "").strip(),
            "brand": str(row[index["MARCA"]] or "").strip(),
            "description": str(row[index["DESCRIPCION"]] or "").strip(),
            "active": bool(row[index["ACTIVO"]]),
            "activeWeb": bool(row[index["ACTIVO WEB"]]),
        }
        products[code] = current
    else:
        current["active"] = current["active"] or bool(row[index["ACTIVO"]])
        current["activeWeb"] = current["activeWeb"] or bool(row[index["ACTIVO WEB"]])

payload = {
    "sourceFilename": SOURCE.name,
    "sheet": sheet.title,
    "codeColumn": "CODIGO",
    "alternativeCodeColumn": "CODIGO ALTERNATIVO",
    "sourceRows": sheet.max_row - 1,
    "products": list(products.values()),
}
OUTPUT.parent.mkdir(parents=True, exist_ok=True)
OUTPUT.write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
print(json.dumps({"sourceRows": payload["sourceRows"], "uniqueProducts": len(products), "outputBytes": OUTPUT.stat().st_size}))
